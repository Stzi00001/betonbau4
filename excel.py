"""Tabellenkalkulationen (Excel mit echten Formeln) zu den Bemessungstools."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

GELB = PatternFill("solid", fgColor="FCE8A6")
GRAU = PatternFill("solid", fgColor="E8EAE9")
DUNKEL = PatternFill("solid", fgColor="172025")
DUENN = Border(bottom=Side(style="thin", color="C9CECC"))


def torsion(pfad):
    wb = Workbook()
    ws = wb.active
    ws.title = "Torsion"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 52

    ws["A1"] = "Torsionsbemessung Rechteckquerschnitt – DIN EN 1992-1-1 + NA (DE)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Gelbe Zellen eingeben. Vorläufig: Annahmen mit dem Vorlesungsskript abgleichen."
    ws["A2"].font = Font(italic=True, color="56625F")

    def kopf(r, text):
        ws.cell(r, 1, text).font = Font(bold=True, color="FFFFFF")
        for c in range(1, 5):
            ws.cell(r, c).fill = DUNKEL

    def zeile(r, name, wert, einheit, info="", eingabe=False, fmt="0.00"):
        ws.cell(r, 1, name)
        z = ws.cell(r, 2, wert)
        z.number_format = fmt
        z.alignment = Alignment(horizontal="right")
        if eingabe:
            z.fill = GELB
        ws.cell(r, 3, einheit)
        ws.cell(r, 4, info).font = Font(color="56625F")
        for c in range(1, 5):
            ws.cell(r, c).border = DUENN

    kopf(4, "Eingabe")
    zeile(5, "Breite b", 40, "cm", eingabe=True, fmt="0.0")
    zeile(6, "Höhe h", 60, "cm", eingabe=True, fmt="0.0")
    zeile(7, "d₁ (Achse Längsbewehrung)", 5, "cm", eingabe=True, fmt="0.0")
    zeile(8, "Betonfestigkeitsklasse", "C30/37", "", "Auswahl", eingabe=True, fmt="@")
    zeile(9, "T_Ed", 40, "kNm", eingabe=True, fmt="0.0")
    zeile(10, "V_Ed", 150, "kN", "0 = reine Torsion", eingabe=True, fmt="0.0")
    zeile(11, "cotθ", 1.0, "–", "1,0 bis 1,75", eingabe=True)
    zeile(12, "Querschnitt", "voll", "", "voll / kasten", eingabe=True, fmt="@")
    zeile(13, "ρ_l", 1.0, "%", "für V_Rd,c", eingabe=True)

    dv = DataValidation(type="list", formula1='"C20/25,C25/30,C30/37,C35/45,C40/50,C45/55,C50/60,C55/67,C60/75"')
    dv2 = DataValidation(type="list", formula1='"voll,kasten"')
    ws.add_data_validation(dv); ws.add_data_validation(dv2)
    dv.add("B8"); dv2.add("B12")

    kopf(15, "Baustoffe")
    zeile(16, "f_ck", '=VALUE(MID(B8,2,FIND("/",B8)-2))', "N/mm²", "aus Festigkeitsklasse", fmt="0")
    zeile(17, "f_cd", "=0.85*B16/1.5", "N/mm²", "α_cc·f_ck/γ_c")
    zeile(18, "f_yd", "=500/1.15", "N/mm²", "f_yk/γ_s", fmt="0.0")
    zeile(19, "ν₂", "=IF(B16>50,1.1-B16/500,1)", "–", "Abminderung hochfester Beton", fmt="0.000")

    kopf(21, "Ersatzhohlkasten")
    zeile(22, "t_ef", "=MIN(2*B7,MIN(B5,B6)/2)", "cm", "2·d₁ (DE NA)", fmt="0.0")
    zeile(23, "A_k", "=(B5-B22)*(B6-B22)", "cm²", "(b − t_ef)·(h − t_ef)", fmt="0")
    zeile(24, "u_k", "=2*((B5-B22)+(B6-B22))", "cm", "Umfang Mittellinie", fmt="0.0")
    zeile(25, "θ", "=DEGREES(ATAN(1/B11))", "°", fmt="0.0")
    zeile(26, "Schubfluss τ_t·t_ef", "=B9*1000000/(2*B23*100)", "N/mm", "T_Ed/(2·A_k)", fmt="0.0")

    kopf(28, "Nachweise")
    zeile(29, "ν", '=IF(B12="kasten",0.75,0.525)*B19', "–", "DE NA", fmt="0.000")
    zeile(30, "T_Rd,max", "=2*B29*B17*B23*100*B22*10*SIN(RADIANS(B25))*COS(RADIANS(B25))/1000000", "kNm",
          "2·ν·f_cd·A_k·t_ef·sinθ·cosθ", fmt="0.0")
    zeile(31, "Ausnutzung T_Ed / T_Rd,max", "=B9/B30", "–", fmt="0.0%")
    zeile(32, "a_sw je Bügelschenkel", "=B9*1000000/(2*B23*100*B18*B11)*10", "cm²/m", "T_Ed/(2·A_k·f_yd·cotθ)")
    zeile(33, "ΣA_sl", "=B9*1000000*B24*10*B11/(2*B23*100*B18)/100", "cm²", "T_Ed·u_k·cotθ/(2·A_k·f_yd)")
    zeile(34, "s_max Bügel", "=MIN(2*(B5+B6)/8,MIN(B5,B6))", "cm", "min(u/8 ; min(b,h))", fmt="0.0")
    zeile(35, "d", "=B6-B7", "cm", fmt="0.0")
    zeile(36, "V_Rd,max", "=0.9*B35*10*B5*10*0.75*B19*B17/(B11+1/B11)/1000", "kN", "b_w·z·ν₁·f_cd/(cotθ+tanθ)", fmt="0.0")
    zeile(37, "Interaktion T + V", '=IF(B12="kasten",B9/B30+B10/B36,(B9/B30)^2+(B10/B36)^2)', "≤ 1",
          "voll: quadratisch, kasten: linear", fmt="0.000")
    zeile(38, "Nachweis erfüllt?", '=IF(AND(B31<=1,B37<=1),"ja","nein")', "", fmt="@")

    kopf(40, "Verzicht auf rechnerischen Nachweis (DE NA 6.3.2(5))")
    zeile(41, "k", "=MIN(1+SQRT(200/(B35*10)),2)", "–", fmt="0.000")
    zeile(42, "V_Rd,c", "=MAX(0.1*B41*(100*MIN(B13/100,0.02)*B16)^(1/3),0.0525/1.5*B41^1.5*SQRT(B16))*B5*10*B35*10/1000",
          "kN", "ohne Querkraftbewehrung", fmt="0.0")
    zeile(43, "Bedingung 1: T ≤ V·b_w/4,5", '=IF(B10>0,IF(B9<=B10*B5/100/4.5,"ja","nein"),"–")', "", fmt="@")
    zeile(44, "Bedingung 2: V·(1+4,5·T/(V·b_w)) ≤ V_Rd,c",
          '=IF(B10>0,IF(B10*(1+4.5*B9/(B10*B5/100))<=B42,"ja","nein"),"–")', "", fmt="@")

    for r in (38, 43, 44):
        ws.cell(r, 2).font = Font(bold=True)
    ws.freeze_panes = "A4"
    wb.save(pfad)


TABELLEN = {"torsion": torsion}
